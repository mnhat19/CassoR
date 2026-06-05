"""
Generate HR data template files:
- static/templates/hr_data_template.xlsx
- static/templates/hr_data_template.csv
"""

import os
import csv
from pathlib import Path

# Try importing openpyxl; if not available, install it
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = BASE_DIR / "static" / "templates"
TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)

HEADERS = [
    "Mã NV",
    "Họ và tên",
    "Chức danh",
    "Nhóm chuyên môn",
    "Phân khúc",
    "Nguồn tuyển",
]

# ---------------------------------------------------------------------------
# 1. XLSX
# ---------------------------------------------------------------------------

def create_xlsx():
    wb = openpyxl.Workbook()

    # ---- Sheet chính ----
    ws_data = wb.active
    ws_data.title = "Dữ liệu nhân sự"

    header_fill = PatternFill(fill_type="solid", fgColor="007E47")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for col_idx, col_name in enumerate(HEADERS, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Auto-fit column widths based on header text length
    for col_idx, col_name in enumerate(HEADERS, start=1):
        col_letter = get_column_letter(col_idx)
        # Give a bit of extra padding
        ws_data.column_dimensions[col_letter].width = max(len(col_name) + 4, 15)

    ws_data.row_dimensions[1].height = 22

    # ---- Sheet hướng dẫn ----
    ws_guide = wb.create_sheet(title="Hướng dẫn")

    title_font = Font(name="Calibri", size=14, bold=True, color="007E47")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    note_font = Font(name="Calibri", size=11, italic=True, color="595959")

    guide_rows = [
        ("HƯỚNG DẪN ĐIỀN DỮ LIỆU NHÂN SỰ", "title"),
        ("", None),
        ("CÁC CỘT DỮ LIỆU", "bold"),
        ("", None),
        ("1. Mã NV  (BẮT BUỘC)", "bold"),
        ("   Mã định danh duy nhất của nhân viên. Ví dụ: EMP001, EMP002", "normal"),
        ("", None),
        ("2. Họ và tên  (không bắt buộc)", "bold"),
        ("   Họ và tên đầy đủ của nhân viên.", "normal"),
        ("", None),
        ("3. Chức danh  (BẮT BUỘC)", "bold"),
        ("   Chức danh kỹ thuật tiếng Anh. Ví dụ hợp lệ:", "normal"),
        ("   • Junior Software Engineer", "normal"),
        ("   • Senior Software Engineer", "normal"),
        ("   • Lead Software Engineer", "normal"),
        ("   • Principal Software Engineer", "normal"),
        ("   • Junior Data Analyst", "normal"),
        ("   • Senior Data Analyst", "normal"),
        ("   • Engineering Manager", "normal"),
        ("", None),
        ("4. Nhóm chuyên môn  (không bắt buộc)", "bold"),
        ("   Tên nhóm chuyên môn. Ví dụ: Engineering, Data, Product", "normal"),
        ("", None),
        ("5. Phân khúc  (không bắt buộc)", "bold"),
        ("   Phân khúc nhân sự. Ví dụ: Technology, Business", "normal"),
        ("", None),
        ("6. Nguồn tuyển  (không bắt buộc)", "bold"),
        ('   Ghi "external" nếu tuyển từ bên ngoài công ty.', "normal"),
        ('   Ghi "internal" nếu chuyển từ bộ phận khác.', "normal"),
        ("   Để trống nếu không rõ.", "normal"),
        ("", None),
        ("LƯU Ý QUAN TRỌNG", "bold"),
        ("   Chỉ cần điền \"Mã NV\" và \"Chức danh\" là đủ để chạy phân tích.", "note"),
        ("   Các cột còn lại là tuỳ chọn, giúp phân tích chi tiết hơn.", "note"),
    ]

    ws_guide.column_dimensions["A"].width = 80

    for row_idx, (text, style) in enumerate(guide_rows, start=1):
        cell = ws_guide.cell(row=row_idx, column=1, value=text)
        if style == "title":
            cell.font = title_font
        elif style == "bold":
            cell.font = bold_font
        elif style == "note":
            cell.font = note_font
        else:
            cell.font = normal_font

    out_path = TEMPLATES_DIR / "hr_data_template.xlsx"
    wb.save(out_path)
    print(f"Created: {out_path}")


# ---------------------------------------------------------------------------
# 2. CSV (UTF-8 BOM)
# ---------------------------------------------------------------------------

def create_csv():
    out_path = TEMPLATES_DIR / "hr_data_template.csv"
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
    print(f"Created: {out_path}")


if __name__ == "__main__":
    create_xlsx()
    create_csv()
    print("Done.")
