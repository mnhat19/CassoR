"""Generate downloadable Excel report from Talent Radar analysis results."""
from __future__ import annotations

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Brand colors
COLOR_HEADER_BG = "007E47"  # Cassor green
COLOR_HEADER_FG = "FFFFFF"
COLOR_ALT_ROW = "F0FBF6"
COLOR_TITLE_BG = "E6F7EF"


def _header_style(ws, row: int, cols: list[str]) -> None:
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.font = Font(bold=True, color=COLOR_HEADER_FG, name="Calibri", size=10)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            bottom=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
        )


def _data_row(ws, row: int, values: list, alt: bool = False) -> None:
    fill = PatternFill("solid", fgColor=COLOR_ALT_ROW) if alt else None
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="center")
        if fill:
            cell.fill = fill
        cell.border = Border(
            bottom=Side(style="hair", color="DDDDDD"),
            right=Side(style="hair", color="DDDDDD"),
        )


def _auto_width(ws, min_width: int = 10, max_width: int = 40) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, max_width))


def _sheet_title(ws, title: str, subtitle: str, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(bold=True, size=14, name="Calibri", color="1A1A1A")
    t.fill = PatternFill("solid", fgColor=COLOR_TITLE_BG)
    t.alignment = Alignment(horizontal="center", vertical="center")
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(size=10, name="Calibri", color="666666", italic=True)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18


def _sheet_summary(ws, analysis: dict) -> None:
    ws.title = "Tổng quan"
    _sheet_title(ws, "Báo cáo Talent Radar", f"Tổng số nhân sự phân tích: {analysis['total_employees']}", 3)
    headers = ["Chỉ số", "Kết quả", "Ghi chú"]
    _header_style(ws, 3, headers)

    rows = [
        ("Tổng nhân sự", analysis["total_employees"], ""),
        ("Số cảnh báo rẽ nhánh", len(analysis["branch_eligible"]), "Professional tại đỉnh track"),
        ("Số rủi ro chững lại", len(analysis["stagnation"]), "Điểm stagnation > 0"),
        ("Số phân khúc phân tích", len(analysis["dependency"]), ""),
    ]

    avg_ext = 0.0
    if analysis["dependency"]:
        avg_ext = sum(d["external_hire_ratio"] for d in analysis["dependency"]) / len(analysis["dependency"])
    rows.append(("Tỷ lệ tuyển ngoài trung bình", f"{avg_ext:.1%}", "Mục tiêu ≤35%"))

    for i, row_data in enumerate(rows):
        _data_row(ws, 4 + i, list(row_data), alt=(i % 2 == 0))

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


def _sheet_pyramid(ws, pyramid: list[dict]) -> None:
    ws.title = "Phân bổ cấp độ"
    _sheet_title(ws, "Kim tự tháp tổ chức", "Phân bổ nhân sự theo chuyên môn và cấp độ", 5)
    headers = ["Mã chuyên môn", "Tên chuyên môn", "Track", "Cấp độ", "Số nhân sự"]
    _header_style(ws, 3, headers)

    sorted_data = sorted(pyramid, key=lambda x: (x["expertise_code"], x["track"], x["level"]))
    for i, entry in enumerate(sorted_data):
        _data_row(ws, 4 + i, [
            entry["expertise_code"],
            entry["expertise_name"],
            entry["track"],
            f"L{entry['level']}",
            entry["count"],
        ], alt=(i % 2 == 0))

    if not sorted_data:
        ws.cell(row=4, column=1, value="Không có dữ liệu phù hợp")

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


def _sheet_stagnation(ws, stagnation: list[dict]) -> None:
    ws.title = "Rủi ro chững lại"
    _sheet_title(ws, "Cảnh báo Stagnation Risk", "Nhân sự có nguy cơ chững lại sự nghiệp", 6)
    headers = ["Mã NV", "Họ và tên", "Chuyên môn", "Cấp độ hiện tại", "Điểm rủi ro", "Mức độ"]
    _header_style(ws, 3, headers)

    sorted_data = sorted(stagnation, key=lambda x: -x["score"])
    for i, entry in enumerate(sorted_data):
        score = entry["score"]
        level_label = "🔴 Nguy hiểm" if score >= 8 else ("🟡 Cao" if score >= 6 else "🟢 Trung bình")
        _data_row(ws, 4 + i, [
            entry["employee_id"],
            entry.get("full_name", ""),
            entry.get("expertise_code", ""),
            f"L{entry['current_level']}",
            score,
            level_label,
        ], alt=(i % 2 == 0))

    if not sorted_data:
        ws.cell(row=4, column=1, value="Không có cảnh báo stagnation")

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


def _sheet_branch(ws, branch_eligible: list[dict]) -> None:
    ws.title = "Đủ điều kiện rẽ nhánh"
    _sheet_title(ws, "Đủ điều kiện rẽ nhánh", "Nhân sự Professional tại đỉnh track, sẵn sàng chuyển hướng", 5)
    headers = ["Mã NV", "Họ và tên", "Mã chuyên môn", "Cấp độ hiện tại", "Khuyến nghị"]
    _header_style(ws, 3, headers)

    for i, entry in enumerate(branch_eligible):
        _data_row(ws, 4 + i, [
            entry["employee_id"],
            entry.get("full_name", ""),
            entry["expertise_code"],
            f"L{entry['current_level']}",
            "Xem xét chuyển sang Management hoặc Leadership track",
        ], alt=(i % 2 == 0))

    if not branch_eligible:
        ws.cell(row=4, column=1, value="Không có nhân sự đủ điều kiện rẽ nhánh")

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


def _sheet_dependency(ws, dependency: list[dict]) -> None:
    ws.title = "Phân tích phụ thuộc"
    _sheet_title(ws, "Phân tích tuyển dụng", "Tỷ lệ tuyển ngoài vs pipeline nội bộ theo phân khúc", 6)
    headers = ["Phân khúc", "Tổng nhân sự", "Tuyển ngoài", "Pipeline nội bộ", "Tỷ lệ tuyển ngoài", "Trạng thái"]
    _header_style(ws, 3, headers)

    for i, entry in enumerate(dependency):
        ratio = entry["external_hire_ratio"]
        status = "✅ Tốt" if ratio <= 0.35 else ("⚠️ Cảnh báo" if ratio <= 0.5 else "❌ Cần cải thiện")
        _data_row(ws, 4 + i, [
            entry["segment"],
            entry["total"],
            entry["external_count"],
            entry["internal_count"],
            f"{ratio:.1%}",
            status,
        ], alt=(i % 2 == 0))

    if not dependency:
        ws.cell(row=4, column=1, value="Không có dữ liệu phân tích")

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


def generate_excel_report(analysis: dict) -> bytes:
    """Generate a styled Excel workbook from radar analysis dict. Returns bytes."""
    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Add sheets
    _sheet_summary(wb.create_sheet(), analysis)
    _sheet_pyramid(wb.create_sheet(), analysis["pyramid"])
    _sheet_stagnation(wb.create_sheet(), analysis["stagnation"])
    _sheet_branch(wb.create_sheet(), analysis["branch_eligible"])
    _sheet_dependency(wb.create_sheet(), analysis["dependency"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
